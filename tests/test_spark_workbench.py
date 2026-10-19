import asyncio
from datetime import datetime, timezone
import io
import json
import math
from pathlib import Path
import sys
import tempfile
import unittest

sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
from spark_workspace import SparkWorkspace
from spark_workbench import Workbench, ACTION_KEYS
from spark_productivity import Productivity, WorkbenchError, transform_text
from spark_robotics import simulate, route
from spark_integrations import Integrations
from spark_workbench_voice import parse_workbench_command


class WorkbenchTests(unittest.IsolatedAsyncioTestCase):
    async def asyncSetUp(self):
        self.directory=tempfile.TemporaryDirectory()
        self.workspace=SparkWorkspace(self.directory.name)
        self.workbench=Workbench(self.workspace)
        self.project='general'

    async def asyncTearDown(self):
        await self.workbench.integrations.close()
        await self.workbench.intelligence.close()
        self.directory.cleanup()

    async def action(self,action,**payload):
        return await self.workbench.execute(action,self.project,payload)

    async def test_voice_status_intents_need_no_task_identifier_and_never_grant_execution(self):
        self.assertEqual(parse_workbench_command('show my agent status')['action'],'agent_snapshot')
        self.assertEqual(parse_workbench_command('check system health')['action'],'intelligence_health')
        self.assertIsNone(parse_workbench_command('approve all agent tasks'))

    async def test_persistent_agent_api_executes_real_status_and_requires_exact_approval(self):
        job=await self.action('agent_create',title='Check core',plan=[{'action':'get_system_status','args':{},'depends':[]}])
        with self.assertRaises(WorkbenchError):await self.action('agent_run',id=job['id'])
        await self.action('agent_authorize',id=job['id'],plan_hash=job['plan_hash'],approved=True)
        result=await self.action('agent_run',id=job['id'])
        self.assertEqual(result['state'],'completed')
        self.assertEqual(result['results'][0]['value']['service'],'SPARK Core')
        snapshot=await self.action('agent_snapshot')
        self.assertEqual(len(snapshot['jobs']),1)
        self.assertTrue(snapshot['traces'][0]['verified'])
        self.assertEqual((await self.workbench.execute('agent_snapshot','private',{}))['jobs'],[])
        await self.workbench.intelligence.close()

    async def test_agent_calls_only_host_approved_mcp_reads_through_real_transport(self):
        hub=Integrations({'test-service':{'command':sys.executable,'args':[str(Path(__file__).parent/'fixtures'/'rea_mcp_server.py')],'read_tools':['read_counter']}})
        self.workbench=Workbench(self.workspace,integrations=hub)
        job=await self.action('agent_create',title='Read counter',plan=[{'action':'mcp_read','args':{'server':'test-service','name':'read_counter','arguments':{}},'depends':[]}])
        await self.action('agent_authorize',id=job['id'],plan_hash=job['plan_hash'],approved=True)
        result=await self.action('agent_run',id=job['id'])
        self.assertEqual(result['state'],'completed')
        self.assertEqual(result['results'][0]['value']['content'][0]['text'],'0')
        denied=await self.action('agent_create',title='Reject counter write',plan=[{'action':'mcp_read','args':{'server':'test-service','name':'increment','arguments':{'amount':1}},'depends':[]}])
        await self.action('agent_authorize',id=denied['id'],plan_hash=denied['plan_hash'],approved=True)
        self.assertEqual((await self.action('agent_run',id=denied['id']))['state'],'failed')

    async def test_memory_health_and_measured_evidence_are_accessible_without_optional_models(self):
        item=await self.action('memory_remember',kind='task',body='Use metric units',source='user',ttl=60)
        result=await self.action('memory_search',query='metric',semantic=True)
        self.assertEqual(result['mode'],'fts5')
        self.assertEqual(result['results'][0]['id'],item['id'])
        health=await self.action('intelligence_health')
        self.assertEqual(health['core']['service'],'SPARK Core')
        self.assertIn('cpu_percent',health['host_vitals'])
        versions=await self.action('intelligence_versions')
        self.assertIsInstance(versions['spark'],str)
        self.assertFalse(versions['automatic_updates'])
        evidence=await self.action('intelligence_evidence')
        self.assertIsNone(evidence['success_rate'])
        self.assertEqual(evidence['verified_steps'],0)
        self.assertIsNone(evidence['energy_wh'])
        await self.workbench.intelligence.close()

    async def test_learning_executes_matched_model_cases_before_activation(self):
        from unittest.mock import AsyncMock
        candidate=await self.action('improvement_propose',content='Answer with the correct arithmetic result only.',kind='prompt')
        self.workbench.intelligence.models.generate=AsyncMock(side_effect=[{'answer':'wrong'},{'answer':'4'}])
        result=await self.action('improvement_evaluate',id=candidate['id'],model='test-local',cases=[{'input':'2+2','expected':'4'}],allow_cloud=False)
        self.assertEqual(result['evidence']['candidate_correct'],1)
        self.assertEqual(self.workbench.intelligence.models.generate.await_count,2)
        activated=await self.action('improvement_activate',id=candidate['id'],approved=True)
        self.assertEqual(activated['status'],'active')
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'done'})
        await self.action('task_model_generate',id='test-local',query='hello',allow_cloud=False)
        self.assertIn('correct arithmetic',self.workbench.intelligence.models.generate.await_args.args[1])
        await self.action('improvement_rollback',id=candidate['id'],approved=True)

    async def test_trace_based_proposals_exclude_private_tool_results_and_stay_staged(self):
        from unittest.mock import AsyncMock
        runtime=self.workbench.intelligence._runtime()
        runtime._save('trace','general','trace-test',{'id':'trace-test','action':'run_skill','verified':True,'failure_class':None,'result':{'private':'private document phrase'},'elapsed':.01})
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'Check arithmetic and cite evidence.'})
        proposed=await self.action('improvement_suggest',model='test-local',kind='few_shot',allow_cloud=False)
        self.assertEqual(proposed['status'],'staged')
        self.assertIn('trace-test',proposed['trace_ids'])
        prompt=self.workbench.intelligence.models.generate.await_args.args[1]
        self.assertNotIn('private document phrase',prompt)
        self.assertIn('verified',prompt)

    async def test_model_usage_history_survives_restart_without_retaining_prompts(self):
        from unittest.mock import AsyncMock
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'private answer','usage':{'input_tokens':3,'output_tokens':1}})
        await self.action('task_model_generate',id='test-local',query='private request',allow_cloud=False)
        await self.workbench.intelligence.close()
        self.workbench=Workbench(SparkWorkspace(self.directory.name))
        result=await self.action('task_models')
        self.assertEqual(len(result['history']),1)
        self.assertEqual(result['history'][0]['input_tokens'],3)
        self.assertNotIn('private',json.dumps(result['history']))

    async def test_model_agent_turns_persist_and_each_tool_action_requires_approval(self):
        from unittest.mock import AsyncMock
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'{"action":"get_system_status","action_input":{}}'})
        first=await self.action('agent_plan',model='test-local',goal='Check status and summarize',allow_cloud=False)
        with self.assertRaises(WorkbenchError):await self.action('agent_continue',session_id=first['session_id'],allow_cloud=False,share_results=True)
        await self.action('agent_authorize',id=first['id'],plan_hash=first['plan_hash'],approved=True)
        await self.action('agent_run',id=first['id'])
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'{"final":"The core is available"}'})
        final=await self.action('agent_continue',session_id=first['session_id'],allow_cloud=False,share_results=True)
        self.assertEqual(final['final'],'The core is available')
        self.assertFalse(final['verified'])
        self.assertEqual(final['session_id'],first['session_id'])

    async def test_instructional_execution_supplies_context_without_granting_tools(self):
        from unittest.mock import AsyncMock
        candidate=await self.action('instructional_stage',content='---\nname: cite-notes\ndescription: Cite sources\n---\nCite supplied source identifiers.',source='inspected SKILL.md')
        await self.action('instructional_test',id=candidate['id'])
        await self.action('instructional_activate',id=candidate['id'],approved=True)
        self.workbench.intelligence.models.generate=AsyncMock(return_value={'answer':'summary','verified':False})
        result=await self.action('task_skill_generate',id='test-local',skill='cite-notes',query='Summarize this sentence.',allow_cloud=False)
        self.assertFalse(result['verified'])
        self.assertIn('Cite supplied',self.workbench.intelligence.models.generate.await_args.args[1])
        self.assertEqual(result['permissions'],[])

    async def test_rag_requires_context_consent_and_never_reads_other_projects(self):
        from unittest.mock import AsyncMock
        self.workspace.execute('import_document','general',{'name':'units.md','data':b'Use metric units in the prototype.'})
        self.workspace.execute('import_document','private',{'name':'secret.md','data':b'Private unrelated confidential phrase.'})
        generate=AsyncMock(return_value={'answer':'Use metric units','verified':False})
        self.workbench.intelligence.models.generate=generate
        with self.assertRaises(WorkbenchError):await self.action('task_rag_generate',id='test-local',query='units',allow_cloud=False,share_context=False)
        generate.assert_not_awaited()
        result=await self.action('task_rag_generate',id='test-local',query='units',allow_cloud=False,share_context=True)
        prompt=generate.await_args.args[1]
        self.assertIn('metric',prompt)
        self.assertNotIn('confidential phrase',prompt)
        self.assertTrue(result['citations'])
        self.assertFalse(result['verified'])

    async def test_retrieval_benchmark_scores_actual_search_against_independent_ids(self):
        document=self.workspace.execute('import_document','general',{'name':'units.md','data':b'Use metric units for engineering.'})
        result=await self.action('memory_benchmark',cases=[{'query':'metric','expected_ids':[document['id']]},{'query':'saturn','expected_ids':[document['id']]}])
        self.assertEqual(result['queries'],2)
        self.assertEqual(result['hit_rate'],.5)
        self.assertEqual(result['mean_reciprocal_rank'],.5)

    async def test_preferences_and_agenda_persist_with_timezone_and_project_isolation(self):
        await self.action('save_preference',key='units',value='metric')
        await self.action('save_preference',key='units',value='SI')
        event=await self.action('add_agenda',title='Review',due='2026-10-19T12:00:00+05:30',kind='reminder')
        self.assertEqual(event['due'],'2026-10-19T06:30:00+00:00')
        restored=Workbench(SparkWorkspace(self.directory.name))
        saved=await restored.execute('workbench_snapshot',self.project,{})
        self.assertEqual(saved['preference'][0]['value'],'SI')
        self.assertEqual(len(saved['preference']),1)
        self.assertEqual((await restored.execute('workbench_snapshot','private',{}))['agenda'],[])
        restored.productivity.clock=lambda:datetime(2026,10,19,6,30,tzinfo=timezone.utc).timestamp()
        self.assertEqual(len(await restored.execute('due_agenda',self.project,{})),1)
        await restored.execute('complete_agenda',self.project,{'id':event['id'],'done':True})
        self.assertEqual(await restored.execute('due_agenda',self.project,{}),[])

    async def test_naive_dates_and_invalid_completion_are_rejected(self):
        for due in ['2026-10-19T12:00:00','bad',None]:
            with self.assertRaises(WorkbenchError):
                await self.action('add_agenda',title='Review',due=due,kind='reminder')
        with self.assertRaises(WorkbenchError):
            await self.action('complete_agenda',id='a'*24,done='true')

    async def test_reviewed_routine_executes_real_skill_and_status_in_order(self):
        skill=self.workspace.skills.propose({'name':'area','description':'Area','operation':'multiply','inputs':['a','b']},self.project)
        self.workspace.skills.test(self.project,skill['id']);self.workspace.skills.activate(self.project,skill['id'])
        routine=await self.action('create_routine',title='Calculate',steps=[{'action':'get_system_status','args':{}},{'action':'run_skill','args':{'name':'area','params':{'a':3,'b':4}}}])
        self.assertEqual(routine['cursor'],0)
        first=await self.action('advance_routine',id=routine['id'],approved=True,step=0)
        self.assertEqual(first['cursor'],1);self.assertEqual(first['status'],'pending')
        with self.assertRaises(WorkbenchError):await self.action('advance_routine',id=routine['id'],approved=True,step=0)
        second=await self.action('advance_routine',id=routine['id'],approved=True,step=1)
        self.assertEqual(second['status'],'completed')
        self.assertIn('12.0',second['steps'][1]['result'])
        self.assertTrue(any('Elapsed' in line for line in second['audit']))

    async def test_denied_or_failed_steps_stop_without_running_later_actions(self):
        calls=[]
        self.workbench.productivity.launcher=lambda app:calls.append(app) or {'ok':False}
        routine=await self.action('create_routine',title='Denied',steps=[{'action':'launch_desktop_app','args':{'app':'notepad'}}])
        denied=await self.action('advance_routine',id=routine['id'],approved=False,step=0)
        self.assertEqual(denied['status'],'cancelled');self.assertEqual(calls,[])
        routine=await self.action('create_routine',title='Failure',steps=[{'action':'launch_desktop_app','args':{'app':'notepad'}},{'action':'launch_desktop_app','args':{'app':'paint'}}])
        failed=await self.action('advance_routine',id=routine['id'],approved=True,step=0)
        self.assertEqual(failed['status'],'failed');self.assertEqual(calls,['notepad'])
        with self.assertRaises(WorkbenchError):await self.action('advance_routine',id=routine['id'],approved=True,step=0)

    async def test_interrupted_action_is_not_replayed_after_restart(self):
        routine=await self.action('create_routine',title='Interrupted',steps=[{'action':'get_system_status','args':{}}])
        routine['status']='running';routine['steps'][0]['status']='running'
        p=self.workbench.productivity
        with p._db(self.project) as db:p._put(db,'routine',routine)
        result=await self.action('advance_routine',id=routine['id'],approved=True,step=0)
        self.assertEqual(result['status'],'failed')
        self.assertEqual(result['steps'][0]['attempts'],0)
        self.assertIn('unknown outcome',' '.join(result['audit']))

    async def test_arbitrary_execution_and_cross_project_access_are_rejected(self):
        for step in [{'action':'run_shell','args':{'command':'whoami'}},{'action':'launch_desktop_app','args':{'app':'powershell'}}]:
            with self.assertRaises(WorkbenchError):await self.action('create_routine',title='Invalid',steps=[step])
        routine=await self.action('create_routine',title='Local',steps=[{'action':'get_system_status','args':{}}])
        with self.assertRaises(WorkbenchError):await self.workbench.execute('advance_routine','other',{'id':routine['id'],'approved':True,'step':0})

    async def test_office_files_round_trip_actual_content_formulas_charts_and_notes(self):
        from docx import Document
        from openpyxl import load_workbook
        from pptx import Presentation
        from pypdf import PdfReader
        spec={'title':'Spark report','sections':[{'heading':'Overview','text':'Original project text','bullets':['Milestone'],'notes':'Speaker note'}]}
        for extension in ['docx','pdf','pptx']:
            result=await self.action('create_office',format=extension,spec=spec)
            blob=(await self.action('download_artifact',id=result['id']))['data']
            if extension=='docx':self.assertIn('Original project text',' '.join(p.text for p in Document(io.BytesIO(blob)).paragraphs))
            elif extension=='pdf':self.assertIn('Original project text',PdfReader(io.BytesIO(blob)).pages[0].extract_text())
            else:
                deck=Presentation(io.BytesIO(blob));self.assertEqual(len(deck.slides),2)
                self.assertIn('Speaker note',deck.slides[1].notes_slide.notes_text_frame.text)
        workbook={'title':'Budget','sheets':[{'name':'Costs','rows':[['Item','Value'],['Materials',10],['Labor',20],['Total',{'formula':'=SUM(B2:B3)'}],['=HYPERLINK("evil")','@literal']], 'chart':True},{'name':'Notes','rows':[['Private data']]}]}
        result=await self.action('create_office',format='xlsx',spec=workbook)
        book=load_workbook(io.BytesIO((await self.action('download_artifact',id=result['id']))['data']))
        self.assertEqual(book.sheetnames,['Costs','Notes'])
        self.assertEqual(book['Costs']['B4'].value,'=SUM(B2:B3)')
        self.assertEqual(book['Costs']['A5'].data_type,'s')
        self.assertEqual(len(book['Costs']._charts),1)
        self.assertEqual(len(await self.action('list_artifacts')),4)
        with self.assertRaises(WorkbenchError):await self.workbench.execute('download_artifact','other',{'id':result['id']})

    async def test_pdf_merge_extract_and_text_read_actual_pages(self):
        spec={'title':'Original PDF','sections':[{'heading':'Title','text':'PDF content for extraction.'}]}
        first=await self.action('create_office',format='pdf',spec=spec)
        blob=(await self.action('download_artifact',id=first['id']))['data']
        imported=await self.action('import_pdf',name='Imported PDF',data=blob)
        merged=await self.action('pdf_operation',ids=[first['id'],imported['id']],pages=[],operation='merge')
        text=await self.action('pdf_operation',ids=[merged['id']],pages=[2],operation='text')
        self.assertEqual(text['pages'],1);self.assertIn('PDF content',text['text'])
        extracted=await self.action('pdf_operation',ids=[merged['id']],pages=[2],operation='extract')
        from pypdf import PdfReader
        self.assertEqual(len(PdfReader(io.BytesIO((await self.action('download_artifact',id=extracted['id']))['data'])).pages),1)
        with self.assertRaises(WorkbenchError):await self.action('pdf_operation',ids=[first['id']],pages=[99],operation='extract')

    async def test_office_validation_never_saves_invalid_spec_or_external_formula(self):
        before=await self.action('list_artifacts')
        for spec in [{'title':'X','sections':[]},{'title':'X','sections':[{'heading':'Test','code':'bad'}]}]:
            with self.assertRaises(WorkbenchError):await self.action('create_office',format='docx',spec=spec)
        with self.assertRaises(WorkbenchError):
            await self.action('create_office',format='xlsx',spec={'title':'X','sheets':[{'name':'Sheet','rows':[[{'formula':'=WEBSERVICE("http://evil")'}]]}]})
        self.assertEqual(await self.action('list_artifacts'),before)


class LabTests(unittest.TestCase):
    def test_kinematics_hits_target_without_hardware(self):
        target=[.3,.1,.2]
        result=simulate({'links':[.3,.25],'steps':[{'type':'move','target':target,'seconds':1},{'type':'gripper','state':'closed'}]})
        for actual,expected in zip(result['frames'][-1]['points'][-1],target):self.assertAlmostEqual(actual,expected,places=8)
        self.assertEqual(result['mode'],'simulation_only');self.assertEqual(result['frames'][-1]['gripper'],'closed')
        self.assertEqual(result['frames'][-1]['time'],1.1)

    def test_simulation_rejects_unreachable_unsafe_types_and_unbounded_sequences(self):
        for target in [[3,0,0],[float('nan'),0,0],[True,0,0]]:
            with self.assertRaises(WorkbenchError):simulate({'links':[.3,.25],'steps':[{'type':'move','target':target,'seconds':1}]})
        with self.assertRaises(WorkbenchError):simulate({'links':[.3,.25],'steps':[{'type':'delay','seconds':10}]*13})

    def test_great_circle_distance_and_antimeridian(self):
        result=route([0,0],[0,90])
        self.assertAlmostEqual(result['distance_km'],10007.557,places=3)
        crossing=route([0,179],[0,-179]);self.assertLess(crossing['distance_km'],223)
        self.assertEqual(route([10,10],[10,10])['distance_km'],0)
        for end in [[91,0],[0,180],[False,0]]:
            with self.assertRaises(WorkbenchError):route([0,0],end)

    def test_clipboard_is_explicit_and_local(self):
        self.assertEqual(transform_text(' {"a":1} ','json')['text'],'{\n  "a": 1\n}')
        self.assertEqual(transform_text(' Hi ','upper')['text'],' HI ')
        self.assertEqual(transform_text('a b\nc','stats'),{'characters':5,'words':3,'lines':2})
        with self.assertRaises(WorkbenchError):transform_text('not JSON','json')

    def test_explicit_voice_intents_have_no_arbitrary_action(self):
        self.assertEqual(parse_workbench_command('Spark, remember my units as metric')['payload'],{'key':'units','value':'metric'})
        now=datetime(2026,10,19,6,0,tzinfo=timezone.utc)
        command=parse_workbench_command('Remind me in 10 minutes to review design',now)
        self.assertEqual(command['payload']['due'],'2026-10-19T06:10:00+00:00')
        self.assertIsNone(parse_workbench_command('Remind me in 99999 hours to wait',now))
        self.assertIsNone(parse_workbench_command('delete my preferences'))


class MCPHubTests(unittest.IsolatedAsyncioTestCase):
    async def test_audit_survives_restart_without_retaining_arguments_or_credentials(self):
        with tempfile.TemporaryDirectory() as directory:
            hub=Integrations({},audit_dir=directory)
            hub._record({'server':'test','action':'mcp_call','ok':False,'arguments':{'secret':'private'},'headers':{'Authorization':'private'}})
            restored=Integrations({},audit_dir=directory)
            history=restored.history()
            self.assertEqual(len(history),1)
            self.assertNotIn('private',json.dumps(history))
            await hub.close();await restored.close()

    async def test_business_adapters_are_disabled_without_host_endpoints_and_authorization(self):
        hub=Integrations({})
        adapters=hub.list()
        self.assertEqual({item['name'] for item in adapters},{'vaos','vyndi-os','vedm','vaos-jarvis','gmail','google-drive','google-calendar','github','slack'})
        self.assertTrue(all(not item['configured'] and item['state']=='unconnected' for item in adapters))
        with self.assertRaises(WorkbenchError):await hub.execute('vaos','mcp_discover',{})
        await hub.close()

    async def test_streamable_http_transport_discovers_and_calls_without_credentials_in_catalog(self):
        import os,socket
        import httpx
        with socket.socket() as reserve:
            reserve.bind(('127.0.0.1',0));port=reserve.getsockname()[1]
        process=await asyncio.create_subprocess_exec(sys.executable,str(Path(__file__).parent/'fixtures'/'rea_mcp_server.py'),env={**os.environ,'SPARK_TEST_MCP_HTTP_PORT':str(port)},stdout=asyncio.subprocess.DEVNULL,stderr=asyncio.subprocess.DEVNULL)
        hub=Integrations({'http-test':{'transport':'streamable-http','url':f'http://127.0.0.1:{port}/mcp','headers':{'X-Test':'private-header'},'read_tools':['read_counter']}})
        try:
            async with httpx.AsyncClient(timeout=1,trust_env=False) as client:
                for _ in range(60):
                    try:await client.get(f'http://127.0.0.1:{port}/mcp');break
                    except httpx.TransportError:await asyncio.sleep(.25)
            self.assertNotIn('private-header',json.dumps(hub.list()))
            tools=await hub.execute('http-test','mcp_discover',{})
            self.assertIn('read_counter',{tool['name'] for tool in tools})
            self.assertEqual((await hub.execute('http-test','mcp_call',{'name':'read_counter','arguments':{}}))['content'][0]['text'],'0')
        finally:
            await hub.close()
            if process.returncode is None:process.terminate()
            await process.wait()

    async def test_host_configured_service_discovers_and_calls_real_mcp_transport(self):
        config={'test-service':{'command':sys.executable,'args':[str(Path(__file__).parent/'fixtures'/'rea_mcp_server.py')],'env':{'TEST_SECRET':'not-exposed'},'read_tools':['read_counter']}}
        hub=Integrations(config)
        try:
            self.assertNotIn('not-exposed',json.dumps(hub.list()))
            tools=await hub.execute('test-service','mcp_discover',{})
            self.assertEqual({t['name'] for t in tools},{'increment','read_counter'})
            with self.assertRaises(WorkbenchError):
                await hub.execute('test-service','mcp_call',{'name':'increment','arguments':{'amount':2}})
            result=await hub.execute('test-service','mcp_call',{'name':'read_counter','arguments':{}})
            self.assertEqual(result['content'][0]['text'],'0')
            previous=hub.clients['test-service']
            with self.assertRaises(WorkbenchError):await hub.restart('test-service',approved=False)
            recovered=await hub.restart('test-service',approved=True)
            self.assertTrue(recovered['verified'])
            self.assertIsNot(previous,hub.clients['test-service'])
            with self.assertRaises(WorkbenchError):await hub.execute('unknown','mcp_discover',{})
        finally:await hub.close()


class WorkbenchSocketTests(unittest.IsolatedAsyncioTestCase):
    async def test_every_intelligence_action_uses_the_existing_owner_and_face_auth_gate(self):
        from tests.test_spark_evolution import EvolutionSocketTests
        from spark_intelligence.service import ACTION_KEYS as INTELLIGENCE_ACTIONS
        from unittest.mock import AsyncMock
        from types import SimpleNamespace
        for sid,authenticated in [('stranger',True),('owner',False),('owner',True)]:
            namespace=EvolutionSocketTests().handler(authenticated=authenticated)
            executor=AsyncMock(return_value={'saved':True})
            namespace['spark_workbench']=SimpleNamespace(execute=executor)
            for action in INTELLIGENCE_ACTIONS:
                with self.subTest(sid=sid,authenticated=authenticated,action=action):
                    executor.reset_mock()
                    await namespace['workspace_action'](sid,{'request_id':'intelligence-guard','action':action,'project':'general','payload':{}})
                    if sid=='owner' and authenticated:executor.assert_awaited_once()
                    else:executor.assert_not_awaited()
                    self.assertEqual(namespace['sio'].emit.await_args.kwargs['room'],sid)

    async def test_office_and_device_actions_remain_owner_bound(self):
        from tests.test_spark_evolution import EvolutionSocketTests
        from unittest.mock import AsyncMock
        from types import SimpleNamespace
        for sid,authenticated in [('stranger',True),('owner',False),('owner',True)]:
            namespace=EvolutionSocketTests().handler(authenticated=authenticated)
            namespace['spark_workbench']=SimpleNamespace(execute=AsyncMock(return_value={'saved':True}))
            await namespace['workspace_action'](sid,{'request_id':'workbench-1','action':'create_office','project':'general','payload':{'format':'docx','spec':{}}})
            if sid=='owner' and authenticated:
                namespace['spark_workbench'].execute.assert_awaited_once()
            else:
                namespace['spark_workbench'].execute.assert_not_awaited()
            self.assertEqual(namespace['sio'].emit.await_args.kwargs['room'],sid)


if __name__=='__main__':unittest.main()
